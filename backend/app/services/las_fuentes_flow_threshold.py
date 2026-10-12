"""LF16: revision optativa de umbrales en un historico individual.

No modifica lecturas ni calidad; usa los mismos buckets que /water/history.
No crea alarmas SCADA ni reglas persistentes.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO
from math import isfinite
from typing import Any



def numeric(value: Any) -> float | None:
    if value is None or value == '':
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if isfinite(number) else None


@dataclass(frozen=True)
class FlowThreshold:
    enabled: bool = False
    direction: str = 'below'
    basis: str = 'extreme'
    min_value: float | None = None
    max_value: float | None = None

    def validate(self) -> None:
        if not self.enabled:
            return
        if self.direction not in {'below', 'above', 'both'}:
            raise ValueError('Sentido invalido: selecciona por debajo, por encima o ambos.')
        if self.basis not in {'extreme', 'average'}:
            raise ValueError('Criterio invalido: selecciona Min/Max o Promedio.')
        if self.direction in {'below', 'both'} and self.min_value is None:
            raise ValueError('Indica el umbral minimo en L/s.')
        if self.direction in {'above', 'both'} and self.max_value is None:
            raise ValueError('Indica el umbral maximo en L/s.')
        if any(value is not None and (not isfinite(value) or value < 0) for value in (self.min_value, self.max_value)):
            raise ValueError('Los umbrales deben ser numeros finitos mayores o iguales a cero.')
        if self.direction == 'both' and self.min_value > self.max_value:
            raise ValueError('El minimo no puede ser mayor que el maximo.')

    def label(self) -> str:
        if not self.enabled:
            return 'Revision desactivada'
        label = 'Promedio' if self.basis == 'average' else 'Min/Max del intervalo'
        if self.direction == 'both':
            return f'Minimo {self.min_value:.2f} / maximo {self.max_value:.2f} L/s - {label}'
        if self.direction == 'below':
            return f'Por debajo de {self.min_value:.2f} L/s - {label}'
        return f'Por encima de {self.max_value:.2f} L/s - {label}'


def evaluate(point: dict[str, Any], config: FlowThreshold) -> str:
    if not config.enabled:
        return 'No revisado'
    samples = numeric(point.get('samples'))
    if samples is None or samples <= 0:
        return 'Sin lectura'
    average = numeric(point.get('flow_avg_lps'))
    minimum = numeric(point.get('flow_min_lps'))
    maximum = numeric(point.get('flow_max_lps'))
    low = average if config.basis == 'average' else (minimum if minimum is not None else average)
    high = average if config.basis == 'average' else (maximum if maximum is not None else average)
    need_low = config.direction in ('below', 'both')
    need_high = config.direction in ('above', 'both')
    if (need_low and low is None) or (need_high and high is None):
        return 'Sin lectura'
    below = need_low and low < config.min_value
    above = need_high and high > config.max_value
    if below and above:
        return 'Fuera de rango'
    if below:
        return 'Bajo umbral'
    if above:
        return 'Sobre umbral'
    return 'Dentro de rango'


INCIDENTS = {'Bajo umbral', 'Sobre umbral', 'Fuera de rango'}


def summarize(points: list[dict[str, Any]], config: FlowThreshold) -> dict[str, int]:
    states = [evaluate(point, config) for point in points]
    return {
        'total': len(states),
        'breached': sum(s in INCIDENTS for s in states),
        'below': sum(s in {'Bajo umbral', 'Fuera de rango'} for s in states),
        'above': sum(s in {'Sobre umbral', 'Fuera de rango'} for s in states),
        'within': states.count('Dentro de rango'),
        'no_reading': states.count('Sin lectura'),
    }


def _extreme(points: list[dict[str, Any]], field: str, first: bool) -> str:
    observed = [(numeric(p.get(field)), str(p.get('bucket_start') or ''))
                for p in points if (numeric(p.get('samples')) or 0) > 0]
    observed = [(value, start) for value, start in observed if value is not None]
    if not observed:
        return 'Sin lectura'
    value, start = (min if first else max)(observed, key=lambda row: row[0])
    return f'{value:.2f} L/s ({start[:16]})'


def diagnosis(points: list[dict[str, Any]], config: FlowThreshold) -> str:
    counts = summarize(points, config)
    parts = [f'{config.label()}. {counts["breached"]} de {counts["total"]} intervalos con incidencia',
             f'{counts["below"]} bajo minimo, {counts["above"]} sobre maximo',
             f'{counts["within"]} dentro de rango, {counts["no_reading"]} sin lectura']
    if config.direction in {'below', 'both'}:
        parts.append('Menor observado: ' + _extreme(points, 'flow_avg_lps' if config.basis == 'average' else 'flow_min_lps', True))
    if config.direction in {'above', 'both'}:
        parts.append('Mayor observado: ' + _extreme(points, 'flow_avg_lps' if config.basis == 'average' else 'flow_max_lps', False))
    return '. '.join(parts) + '. Los conteos por limite pueden coincidir en un intervalo. Las horas corresponden a los intervalos, no al minuto exacto del evento.'


def _volume_rows(points: list[dict[str, Any]], mode: str):
    cumulative = 0.0
    for point in points:
        volume = numeric(point.get('volume_m3')) if point.get('volume_reliable') is True else None
        if volume is not None:
            cumulative += volume
        yield point, (volume if mode == 'interval' else (cumulative if volume is not None else None))


def _fetch(module: str, sensor_id: int, start_date: str, end_date: str, aggregation: str, force_refresh: bool):
    from app.services.insurgentes_history_service import get_insurgentes_water_history
    payload = get_insurgentes_water_history(module=module, sensor_id=sensor_id, start_date=start_date,
                                            end_date=end_date, aggregation=aggregation, force_refresh=force_refresh)
    if not any((numeric(p.get('samples')) or 0) > 0 for p in payload.get('points') or []):
        raise ValueError('No hay lecturas historicas para este sensor y periodo.')
    return payload


def export_detail_pdf(*, module: str, sensor_id: int, start_date: str, end_date: str,
                      aggregation: str, detail_volume_display: str = 'interval',
                      threshold: FlowThreshold, force_refresh: bool = False) -> tuple[bytes, str]:
    threshold.validate()
    if detail_volume_display not in {'interval', 'cumulative'}:
        raise ValueError('Modo de volumen invalido.')
    payload = _fetch(module, sensor_id, start_date, end_date, aggregation, force_refresh)
    points = list(payload['points'])

    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from app.services.report_visual_theme import pdf_footer, pdf_table

    fig, axis = plt.subplots(figsize=(10.7, 3.8), dpi=110)
    volume_axis = axis.twinx()
    timestamps = [datetime.fromisoformat(str(p['bucket_start'])) for p in points]
    flows = [numeric(p.get('flow_avg_lps')) for p in points]
    volumes = [amount for _, amount in _volume_rows(points, detail_volume_display)]
    axis.plot(timestamps, [v if v is not None else float('nan') for v in flows], color='#0EA5E9', linewidth=1.7, label='Flujo promedio (L/s)')
    volume_axis.plot(timestamps, [v if v is not None else float('nan') for v in volumes], color='#A855F7', linewidth=1.3,
                     label=('Volumen acumulado (m3)' if detail_volume_display == 'cumulative' else 'Volumen por intervalo (m3)'))
    if threshold.enabled:
        if threshold.direction in ('below', 'both'):
            axis.axhline(threshold.min_value, color='#D97706', linestyle='--', linewidth=1.2, label=f'Minimo {threshold.min_value:.2f} L/s')
        if threshold.direction in ('above', 'both'):
            axis.axhline(threshold.max_value, color='#E11D48', linestyle='--', linewidth=1.2, label=f'Maximo {threshold.max_value:.2f} L/s')
    if timestamps:
        axis.set_xlim(min(timestamps), max(timestamps) if len(set(timestamps)) > 1 else timestamps[0] + timedelta(minutes=1))
    axis.xaxis.set_major_formatter(mdates.DateFormatter('%d/%m %H:%M'))
    axis.grid(alpha=0.17, linestyle='--')
    axis.set_ylabel('Flujo (L/s)')
    volume_axis.set_ylabel('Volumen (m3)')
    fig.autofmt_xdate(rotation=28)
    handles, labels = axis.get_legend_handles_labels()
    handles2, labels2 = volume_axis.get_legend_handles_labels()
    fig.legend(handles + handles2, labels + labels2, loc='upper center', bbox_to_anchor=(0.5, 1.0), ncol=2, fontsize=8)
    fig.tight_layout(rect=(0, 0, 1, 0.9))
    chart = BytesIO()
    fig.savefig(chart, format='png', dpi=110, facecolor='white')
    plt.close(fig)
    chart.seek(0)

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm,
                            topMargin=12*mm, bottomMargin=12*mm, title='Historico individual ARCA Las Fuentes')
    styles = getSampleStyleSheet()
    heading = ParagraphStyle('LfThresholdTitle', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=15, textColor=colors.HexColor('#163344'))
    body = ParagraphStyle('LfThresholdBody', parent=styles['BodyText'], fontSize=8.7, leading=12, spaceAfter=3*mm)
    from xml.sax.saxutils import escape
    safe = lambda value: escape(str(value))
    story = [Paragraph(f'ARCA Las Fuentes - {safe(payload["name"])}', heading),
             Paragraph(f'Sensor {sensor_id} | Periodo {start_date} a {end_date} | Agrupacion {aggregation} | Volumen {detail_volume_display}', body),
             Image(chart, width=252*mm, height=94*mm), Spacer(1, 4*mm)]
    if threshold.enabled:
        story.append(Paragraph('<b>Diagnostico de umbrales:</b> ' + safe(diagnosis(points, threshold)), body))
    nonempty = [p for p in points if (numeric(p.get('samples')) or 0) > 0]
    average = [numeric(p.get('flow_avg_lps')) for p in nonempty]
    average = [v for v in average if v is not None]
    total_volume = sum(numeric(p.get('volume_m3')) or 0 for p in nonempty if p.get('volume_reliable') is True)
    story.append(Paragraph(f'<b>Promedio por intervalo:</b> {sum(average)/len(average):.2f} L/s' if average else '<b>Promedio:</b> sin lectura', body))
    story.append(Paragraph(f'<b>Volumen validado:</b> {total_volume:.2f} m3 | <b>Totalizador inicial / final:</b> '
                           f'{_fmt(nonempty[0].get("totalizer_open_m3"))} / {_fmt(nonempty[-1].get("effective_totalizer_close_m3") if nonempty[-1].get("effective_totalizer_close_m3") is not None else nonempty[-1].get("totalizer_close_m3"))} m3', body))
    cols = ['Inicio', 'Fin', 'Prom. L/s', 'Min. L/s', 'Max. L/s',
            ('Vol. acum. m3' if detail_volume_display == 'cumulative' else 'Vol. interv. m3'),
            'Tot. m3', 'Muestras', 'Calidad', 'Umbral']
    rows = [cols]
    marked = []
    for point, displayed_volume in _volume_rows(points, detail_volume_display):
        index = len(rows)
        status = evaluate(point, threshold)
        if status in INCIDENTS:
            marked.append(index)
        rows.append([str(point.get('bucket_start') or '')[:16], str(point.get('bucket_end') or '')[:16],
                     _fmt(point.get('flow_avg_lps')), _fmt(point.get('flow_min_lps')), _fmt(point.get('flow_max_lps')),
                     _fmt(displayed_volume), _fmt(point.get('effective_totalizer_close_m3') if point.get('effective_totalizer_close_m3') is not None else point.get('totalizer_close_m3')),
                     str(point.get('samples') or 0), str(point.get('quality_label') or point.get('quality_status') or point.get('data_status') or 'Sin datos'), status])
    table = pdf_table(rows, [26*mm, 26*mm, 19*mm, 19*mm, 19*mm, 22*mm, 21*mm, 15*mm, 24*mm, 28*mm], font_size=6.7, max_width=doc.width)
    if marked:
        table.setStyle(TableStyle([('BACKGROUND', (0, index), (-1, index), colors.HexColor('#FFF1CC')) for index in marked]))
    story.append(table)
    doc.build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)
    return output.getvalue(), f'ARCA_Las_Fuentes_umbral_{sensor_id}_{start_date}_{end_date}.pdf'


def _fmt(value: Any) -> str:
    number = numeric(value)
    return f'{number:.2f}' if number is not None else '-'


def export_detail_excel(*, module: str, sensor_id: int, start_date: str, end_date: str,
                        aggregation: str, detail_volume_display: str = 'interval',
                        threshold: FlowThreshold, force_refresh: bool = False) -> tuple[bytes, str]:
    threshold.validate()
    if detail_volume_display not in {'interval', 'cumulative'}:
        raise ValueError('Modo de volumen invalido.')
    payload = _fetch(module, sensor_id, start_date, end_date, aggregation, force_refresh)
    points = list(payload['points'])
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from app.services.report_visual_theme import style_excel_table

    book = Workbook()
    sheet = book.active
    sheet.title = 'Historico individual'
    sheet.append([f'ARCA Las Fuentes - {payload["name"]} (sensor {sensor_id})'])
    sheet.append([f'Periodo: {start_date} a {end_date} | Agrupacion: {aggregation} | Volumen: {detail_volume_display}'])
    sheet.append(['Diagnostico de umbrales', diagnosis(points, threshold) if threshold.enabled else 'Revision desactivada'])
    sheet.merge_cells('B3:J3')
    sheet['B3'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[3].height = 45
    sheet.append([])
    header = ['Inicio intervalo', 'Fin intervalo', 'Flujo promedio (L/s)', 'Flujo minimo (L/s)',
              'Flujo maximo (L/s)', ('Volumen acumulado (m3)' if detail_volume_display == 'cumulative' else 'Volumen intervalo (m3)'),
              'Totalizador cierre (m3)', 'Muestras', 'Calidad', 'Umbral']
    sheet.append(header)
    for point, displayed_volume in _volume_rows(points, detail_volume_display):
        totalizer = point.get('effective_totalizer_close_m3')
        if totalizer is None:
            totalizer = point.get('totalizer_close_m3')
        sheet.append([str(point.get('bucket_start') or ''), str(point.get('bucket_end') or ''),
                      numeric(point.get('flow_avg_lps')), numeric(point.get('flow_min_lps')),
                      numeric(point.get('flow_max_lps')), displayed_volume, numeric(totalizer),
                      int(point.get('samples') or 0),
                      str(point.get('quality_label') or point.get('quality_status') or point.get('data_status') or 'Sin datos'),
                      evaluate(point, threshold)])
    for col, width in enumerate([22, 22, 22, 22, 22, 25, 24, 12, 24, 24], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    style_excel_table(sheet, 5, sheet.max_row, 10)
    sheet['A1'].font = Font(name='Calibri', size=16, bold=True, color='163344')
    sheet['A3'].font = Font(bold=True, color='096D94')
    warning_fill = PatternFill(fill_type='solid', fgColor='FFF1CC')
    for row in sheet.iter_rows(min_row=6, max_row=sheet.max_row):
        if row[-1].value in INCIDENTS:
            for cell in row:
                cell.fill = warning_fill
    output = BytesIO()
    book.save(output)
    return output.getvalue(), f'ARCA_Las_Fuentes_umbral_{sensor_id}_{start_date}_{end_date}.xlsx'
