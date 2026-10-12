"""Pruebas LF16 sin SQL Server: motor, cero real, Excel y PDF."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services import las_fuentes_flow_threshold as mod


def _point(start, avg, min_value, max_value, samples=15, volume=4.5):
    return {
        'bucket_start': f'2026-10-09T{start}:00', 'bucket_end': '2026-10-09T02:30:00',
        'samples': samples, 'flow_avg_lps': avg, 'flow_min_lps': min_value,
        'flow_max_lps': max_value, 'volume_m3': volume, 'volume_reliable': volume is not None,
        'totalizer_open_m3': 100.0, 'totalizer_close_m3': 104.5, 'sensor_id': 1001,
    }


def test_evaluation_extreme_vs_average():
    point = _point('02:15', 8.0, 0.5, 16.2)
    extreme = mod.FlowThreshold(True, 'both', 'extreme', 1.0, 15.0)
    avg = mod.FlowThreshold(True, 'both', 'average', 1.0, 15.0)
    assert mod.evaluate(point, extreme) == 'Fuera de rango'
    assert mod.evaluate(point, avg) == 'Dentro de rango'
    assert mod.evaluate(point, mod.FlowThreshold(True, 'below', 'extreme', 0.5)) == 'Dentro de rango'
    assert mod.evaluate(_point('02:00', 0.0, 0.0, 0.0), mod.FlowThreshold(True, 'below', 'average', 1.0)) == 'Bajo umbral'
    assert mod.evaluate(_point('02:00', None, None, None, 0, None), extreme) == 'Sin lectura'
    assert mod.summarize([point, _point('02:00', None, None, None, 0, None)], extreme) == {
        'total': 2, 'breached': 1, 'below': 1, 'above': 1, 'within': 0, 'no_reading': 1,
    }


@pytest.mark.parametrize('config', [
    mod.FlowThreshold(True, 'both', 'extreme', 20, 1),
    mod.FlowThreshold(True, 'below', 'extreme'),
    mod.FlowThreshold(True, 'above', 'extreme', None, -1),
    mod.FlowThreshold(True, 'both', 'average', float('nan'), 2),
])
def test_invalid(config):
    with pytest.raises(ValueError):
        config.validate()


def test_report_exports_share_diagnosis(monkeypatch):
    points = [_point('02:00', 0, 0, 0), _point('02:15', 8, 0.5, 16.2),
              _point('02:30', None, None, None, 0, None)]
    monkeypatch.setattr(mod, '_fetch', lambda *args, **kwargs: {
        'name': 'Pozo de prueba', 'points': points,
    })
    cfg = mod.FlowThreshold(True, 'both', 'extreme', 1, 15)
    options = dict(module='well', sensor_id=1001, start_date='2026-10-09',
                   end_date='2026-10-09', aggregation='quarter_hour', threshold=cfg)
    pdf, pdf_name = mod.export_detail_pdf(**options)
    assert pdf.startswith(b'%PDF') and pdf_name.endswith('.pdf')
    xlsx, xlsx_name = mod.export_detail_excel(**options)
    assert xlsx.startswith(b'PK') and xlsx_name.endswith('.xlsx')
    sheet = load_workbook(BytesIO(xlsx)).active
    assert '2 de 3 intervalos' in sheet['B3'].value
    assert [sheet.cell(row, 10).value for row in (6, 7, 8)] == [
        'Bajo umbral', 'Fuera de rango', 'Sin lectura',
    ]
    assert sheet['F8'].value is None
    assert sheet['A6'].fill.fgColor.rgb[-6:] == 'FFF1CC'
    assert sheet['A7'].fill.fgColor.rgb[-6:] == 'FFF1CC'
    assert sheet['A8'].fill.fgColor.rgb[-6:] != 'FFF1CC'
